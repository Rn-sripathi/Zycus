"""Turn an uploaded file into contract text the segmenter can read.

Deterministic: no model involved, so a mangled upload is a parsing bug rather
than an unpredictable one.

The awkward part is not reading the bytes, it is line breaks. The segmenter finds
clauses by looking for "1." at the start of a line, and PDF extraction frequently
returns a paragraph as one long line with the numbering buried mid-sentence. So
extraction is always followed by normalisation that puts clause numbers back at
the start of their own line.
"""

from __future__ import annotations

import io
import logging
import re

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = (".txt", ".md", ".text", ".pdf", ".docx")

MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10 MB


class UnsupportedDocument(ValueError):
    """The file type is not one we can read."""


class UnreadableDocument(ValueError):
    """The file is the right type but could not be parsed."""


def extract_text(filename: str, data: bytes) -> str:
    """Return contract text from an uploaded file.

    Raises UnsupportedDocument or UnreadableDocument with a message meant for
    the person who uploaded the file, not for a log.
    """
    if not data:
        raise UnreadableDocument("That file is empty.")

    if len(data) > MAX_UPLOAD_BYTES:
        raise UnreadableDocument(
            f"That file is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB."
        )

    suffix = _suffix(filename)

    if suffix in (".txt", ".md", ".text"):
        text = _from_plain_text(data)
    elif suffix == ".pdf":
        text = _from_pdf(data)
    elif suffix == ".docx":
        text = _from_docx(data)
    else:
        raise UnsupportedDocument(
            f"Cannot read {suffix or 'that file type'}. "
            f"Supported: {', '.join(SUPPORTED_EXTENSIONS)}."
        )

    text = normalise(text)

    if not text.strip():
        raise UnreadableDocument(
            "No text could be read from that file. If it is a scanned document, "
            "it needs to be run through OCR first."
        )

    return text


def normalise(text: str) -> str:
    """Put numbered clauses back at the start of their own lines.

    PDF extraction in particular returns runs like
    "...written notice. 2. Fees and Payment. Customer shall..." which the
    segmenter cannot see, because it anchors on line starts.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Join words broken across lines by a hyphen ("indemn-\nification").
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)

    # A clause number that follows the end of a sentence starts a new line.
    text = re.sub(r"(?<=[.;:\"')\]])[ \t]+(?=\d{1,2}[.)]\s+[A-Z])", "\n\n", text)

    # Collapse runs of spaces, but never across line breaks.
    text = re.sub(r"[ \t]{2,}", " ", text)

    # At most one blank line between paragraphs.
    text = re.sub(r"\n{3,}", "\n\n", text)

    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def _suffix(filename: str) -> str:
    _, _, ext = (filename or "").rpartition(".")
    return f".{ext.lower()}" if ext else ""


def _from_plain_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnreadableDocument("That text file uses an encoding we cannot read.")


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnreadableDocument("That PDF is password protected.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except UnreadableDocument:
        raise
    except Exception as exc:  # noqa: BLE001 - surfaced to the uploader
        logger.warning("PDF parse failed: %s", exc)
        raise UnreadableDocument("That PDF could not be read.") from exc

    return "\n\n".join(page for page in pages if page.strip())


def _from_docx(data: bytes) -> str:
    try:
        import docx

        document = docx.Document(io.BytesIO(data))
        blocks = [p.text for p in document.paragraphs]

        # Contracts often keep clauses in tables; ignoring them loses the contract.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(" ".join(cells))
    except Exception as exc:  # noqa: BLE001 - surfaced to the uploader
        logger.warning("DOCX parse failed: %s", exc)
        raise UnreadableDocument("That Word document could not be read.") from exc

    return "\n".join(block for block in blocks if block.strip())
